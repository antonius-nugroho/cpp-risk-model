"""Run the calibrated risk model and write the results workbook and charts.

    python calibrate.py        # (re)build config/model_config.yaml from the data files
    python run.py              # simulate, analyse, report
    python run.py -n 5000      # quicker run
"""
from __future__ import annotations

import argparse
import time
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import pandas as pd
import yaml
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import reporting as rp
from model import simulate_plant

FONT = "Arial"


def parse_args():
    ap = argparse.ArgumentParser(description="Calibrated power plant risk model")
    ap.add_argument("--config", default="config/model_config.yaml")
    ap.add_argument("--evidence", default="config/calibration_evidence.xlsx")
    ap.add_argument("-n", "--iterations", type=int)
    ap.add_argument("--seed", type=int)
    ap.add_argument("--out", default="outputs")
    ap.add_argument("--no-backtest", action="store_true")
    ap.add_argument("--no-scenarios", action="store_true")
    return ap.parse_args()


def main():
    a = parse_args()
    cfg = yaml.safe_load(Path(a.config).read_text(encoding="utf-8"))
    n = a.iterations or int(cfg["simulation"]["iterations"])
    seed = a.seed if a.seed is not None else int(cfg["simulation"]["seed"])
    out = Path(a.out)
    out.mkdir(parents=True, exist_ok=True)
    evidence = pd.read_excel(a.evidence, sheet_name=None)
    t0 = time.time()
    print(f"{cfg['name']} | forecast year {cfg['forecast_year']} | {n:,} iterations")

    res = simulate_plant(cfg, n, seed)
    plan = simulate_plant(cfg, 1, seed, plan=True)
    summary = rp.summary_table(res, plan)
    targets = rp.target_table(cfg, res)
    register = rp.risk_register(cfg, res)
    bridge = rp.loss_bridge(cfg, res, plan)
    sens_prod = rp.sensitivity(res, "net_gwh")
    sens_coal = rp.sensitivity(res, "coal_kt")
    bt = pd.DataFrame() if a.no_backtest else rp.backtest(cfg, evidence, n, seed)
    if a.no_scenarios:
        scen, scen_res = pd.DataFrame(), {}
    else:
        scen, scen_res = rp.scenarios(cfg, n, seed, res, simulate_plant)

    pd.set_option("display.width", 200, "display.max_columns", 12, "display.float_format", lambda v: f"{v:,.2f}")
    key = summary[summary["Metric"].isin(["EAF (%)", "EFOR (%)", "Gross production (GWh)", "Net production / sales (GWh)",
                                          "Coal consumption (kt)"])]
    print("\nKEY RESULTS\n", key.drop(columns="Std dev").to_string(index=False))
    print("\nTARGETS\n", targets.to_string(index=False, formatters={"Probability of meeting": "{:.0%}".format}))
    print("\nTOP RISKS\n", register[["Rank", "Event class", "Events per year (plant)", "Expected net energy loss (GWh/yr)",
                                     "Expected EAF impact (pp, plant)"]].head(10).to_string(index=False))
    if not bt.empty:
        print(f"\nBACKTEST: {bt['Inside P10-P90'].sum()} of {len(bt)} actual values inside the simulated P10-P90 band")
    if not scen.empty:
        print("\nSCENARIOS\n", scen[["Scenario", "Mean net GWh", "Delta mean net GWh vs base", "Delta P10 net GWh vs base",
                                     "Mean coal (kt)"]].to_string(index=False))

    charts = [
        rp.chart_distributions(cfg, res, out),
        rp.chart_risk_matrix(register, out),
        rp.chart_pareto(register, out),
        rp.chart_bridge(bridge, out),
        rp.chart_tornado(sens_prod, out, "Drivers of plant net production", "05_tornado_production.png"),
    ]
    if not bt.empty:
        charts.append(rp.chart_backtest(bt, out))
    if scen_res:
        charts.append(rp.chart_scenarios(scen_res, out))

    xl = out / "Risk_Model_Results.xlsx"
    write_workbook(xl, cfg, n, seed, summary, targets, register, bridge, sens_prod, sens_coal, bt, scen, evidence, charts)
    print(f"\nWritten {xl} and {len(charts)} charts in {time.time() - t0:.0f} s")


# --------------------------------------------------------------------------
def write_workbook(path, cfg, n, seed, summary, targets, register, bridge, sens_prod, sens_coal, bt, scen, ev, charts):
    readme = pd.DataFrame({"Item": [
        "Model", "Forecast year", "Iterations / seed", "Data", "Event model", "Duration model", "Availability",
        "Production", "Fuel", "Planned outage", "Status mapping", "Backtest", "PLACEHOLDERS",
        "Values", "Risk score"],
        "Description": [
        cfg["name"],
        str(cfg["forecast_year"]),
        f"{n:,} / {seed}",
        "Failure_Data_highlighted.xlsx (events) and Data_Pengusahaan.xlsx (monthly production, 2023-2025)",
        "Poisson frequency per event class, pooled over both units; rate uncertainty ~ Gamma(k+0.5, 1/unit-years). "
        "Outage extensions: probability per inspection type.",
        "Empirical resampling for classes with >= 8 events; lognormal matched to the observed mean otherwise.",
        "EAF = (PH - POH - SEH - MOH - FOH - EFDH - EMDH - EPDH)/PH; derating eq. hours = loss MWh / 100 MW.",
        "Net = DMN 85 MW x EAF x PH x net output factor; gross = net / (1 - aux share).",
        "Coal = net x NPHR x (1 - biomass share - HSD share) / coal GCV.",
        "Triangular from historical years of the same inspection type (extensions excluded).",
        "Only statuses marked Include = Yes in the Status Mapping sheet of Failure_Data_highlighted.xlsx.",
        "In-sample: rates were estimated from the same years, so it checks consistency, not predictive skill.",
        "EAF/EFOR targets, coal and biomass prices - replace with actual values in config/model_config.yaml.",
        "All numbers in this workbook are simulation outputs written as values; rerun run.py to update them.",
        "Risk register: score = Likelihood x Consequence (formula). Rating: 1-4 Low, 5-9 Medium, 10-14 High, 15-25 Extreme.",
    ]})
    sheets = {
        "Read Me": readme, "Summary": summary, "Targets": targets, "Risk Register": register,
        "Production Bridge": bridge, "Sensitivity - Production": sens_prod, "Sensitivity - Coal": sens_coal,
        "Scenarios": scen, "Backtest": bt,
        "Calib - Event Classes": ev["event_classes"], "Calib - Annual Data": ev["annual_stats"],
        "Calib - Continuous": ev["continuous_fits"], "Calib - Planned Outage": ev["planned_outage"],
        "Calib - Correlations": ev["correlations"], "Events (merged)": ev["events_merged"],
    }
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, t in sheets.items():
            if t is not None and not t.empty:
                t.to_excel(xw, sheet_name=name, index=False)
    wb = load_workbook(path)
    hdr_fill = PatternFill("solid", start_color="1F4E78")
    thin = Side(style="thin", color="D9D9D9")
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                c.font = Font(name=FONT, size=9, bold=c.row == 1, color="FFFFFF" if c.row == 1 else "000000")
                c.border = Border(bottom=thin)
                if c.row == 1:
                    c.fill = hdr_fill
                    c.alignment = Alignment(wrap_text=True, vertical="center")
                elif isinstance(c.value, float):
                    head = str(ws.cell(1, c.column).value)
                    c.number_format = "0.0%" if ("P(" in head or "Share" in head or "Probability" in head
                                                  or "Percentile" in head) else "#,##0.00"
        for col in ws.columns:
            w = max(len(str(c.value)) if c.value is not None else 0 for c in col[:200])
            ws.column_dimensions[col[0].column_letter].width = min(max(9, w * 0.9 + 2), 55)
        ws.row_dimensions[1].height = 32
        ws.freeze_panes = "B2"
    wb["Read Me"].column_dimensions["B"].width = 120

    # risk score and rating as formulas
    ws = wb["Risk Register"]
    heads = {ws.cell(1, j).value: j for j in range(1, ws.max_column + 1)}
    lc, cc = get_column_letter(heads["Likelihood score (1-5)"]), get_column_letter(heads["Consequence score (1-5)"])
    sc, rc = ws.max_column + 1, ws.max_column + 2
    ws.cell(1, sc, "Risk score (L x C)")
    ws.cell(1, rc, "Rating")
    for j in (sc, rc):
        ws.cell(1, j).font = Font(name=FONT, size=9, bold=True, color="FFFFFF")
        ws.cell(1, j).fill = hdr_fill
        ws.column_dimensions[get_column_letter(j)].width = 12
    s_col = get_column_letter(sc)
    for i in range(2, ws.max_row + 1):
        ws.cell(i, sc, f"={lc}{i}*{cc}{i}").font = Font(name=FONT, size=9)
        ws.cell(i, rc, f'=IF({s_col}{i}>=15,"Extreme",IF({s_col}{i}>=10,"High",IF({s_col}{i}>=5,"Medium","Low")))').font = Font(name=FONT, size=9)
    rng = f"{get_column_letter(rc)}2:{get_column_letter(rc)}{ws.max_row}"
    for txt, color in [("Extreme", "F8696B"), ("High", "F8A870"), ("Medium", "FFEB84"), ("Low", "C6E0B4")]:
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{txt}"'], fill=PatternFill("solid", start_color=color)))
    ws.column_dimensions["D"].width = 45

    cs = wb.create_sheet("Charts", 1)
    r = 1
    for p in charts:
        img = XLImage(str(p))
        img.width, img.height = img.width * 0.55, img.height * 0.55
        cs.add_image(img, f"A{r}")
        r += int(img.height / 20) + 2
    wb.save(path)


if __name__ == "__main__":
    main()
