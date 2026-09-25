"""Excel templates for the three uploads (failure log, Data Pengusahaan, pricing).

Each template has an empty data sheet first (the sheet the loaders read), a "Columns" sheet explaining every
column with an example value, and - for the failure log - a "Status Mapping" sheet in the layout load_mapping reads.
The column lists are kept in line with calibrate.FAILURE_COLUMNS / PRODUCTION_COLUMNS / PRICING_COLUMNS.

    python -m app_lib.templates [out_dir]     # writes the three templates to templates/
"""
from __future__ import annotations

import io
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import calibrate as cal
from app_lib.status_codes import DEFAULT_INCLUDED, STATUS_CODES, default_category

FONT = "Arial"
REQ_FILL = PatternFill("solid", fgColor="1F3A5F")
OPT_FILL = PatternFill("solid", fgColor="7F8C9A")
EDIT_FILL = PatternFill("solid", fgColor="FFF2CC")
DATA_ROWS = 2000  # rows pre-formatted for dates / numbers

# (column, unit, description, example, kind) - kind: date | datetime | number | text
FAILURE = [
    ("timestamp_start", "date-time", "Start of the outage or derating", "2025-03-14 08:30", "datetime"),
    ("timestamp_stop", "date-time", "End of the outage or derating", "2025-03-15 02:10", "datetime"),
    ("unit_no", "text", "Unit name - must be written exactly as the 'unit' column in Data Pengusahaan", "Unit 1", "text"),
    ("unit_status", "code", "Status code, e.g. FO, MO, FD, MD, PD, PO (see the Status Mapping sheet)", "FO", "text"),
    ("failure_cause_code", "text", "Cause code as '<code> <component> - <system> - <area>'. The system part groups "
     "events into failure modes (tube leak, coal feeding, fans, ...)", "10 Waterwall - Boiler Tube Leaks - Boiler", "text"),
    ("failure_impact", "text", "Stopped, Tripped or Derating", "Tripped", "text"),
    ("loss_output_mw", "MW", "Gross MW lost (full unit capacity for an outage)", 100, "number"),
    ("failure_duration_hours", "hours", "Duration of the event", 17.67, "number"),
    ("loss_output_mwh", "MWh", "Energy lost = loss_output_mw x duration", 1767, "number"),
    ("root_cause_failure_analysis", "text", "Description / root cause. For PO rows write 'Serious Inspection' "
     "to mark a serious (major) inspection year", "Waterwall tube leak at elevation 32 m", "text"),
]
FAILURE_OPTIONAL = [
    ("power_gross_realization", "MW", "Gross MW actually produced during the event; with loss_output_mw it tells "
     "the model the derating reference MW (default 100 MW if absent)", 0, "number"),
]

PRODUCTION = [
    ("end_of_month", "date", "Last day of the reporting month (one row per unit per month)", "2025-01-31", "date"),
    ("unit", "text", "Unit name - same spelling as unit_no in the failure log and unit_name in Pricing", "Unit 1", "text"),
    ("kwh_produksi_kwh", "kWh", "Gross generation", 62000000, "number"),
    ("kwh_netto_penjualan_kwh", "kWh", "Net generation / sales", 55000000, "number"),
    ("ph_periodehours_jam", "hours", "Period hours in the month", 744, "number"),
    ("sh_servicehours_jam", "hours", "Service hours (on line)", 700, "number"),
    ("ah_availablehours_jam", "hours", "Available hours", 720, "number"),
    ("foh_forcedoutagehours_jam", "hours", "Forced outage hours", 24, "number"),
    ("poh_plannedoutagehours_jam", "hours", "Planned outage hours", 0, "number"),
    ("moh_maintenanceoutagehours_jam", "hours", "Maintenance outage hours", 0, "number"),
    ("fo_omc_forcedoutageomchours_jam", "hours", "Forced outage hours outside management control (OMC)", 0, "number"),
    ("efdh_equivalentforcedderatinghours_jam", "hours", "Equivalent forced derating hours", 12.5, "number"),
    ("emdh_equivalentmaintenancederatinghours_jam", "hours", "Equivalent maintenance derating hours", 0, "number"),
    ("epdh_eqplannedderatinghours_jam", "hours", "Equivalent planned derating hours", 0, "number"),
    ("ncf_netcapacityfactor_pct", "%", "Net capacity factor, in percent (e.g. 75.3) - used to detect the unit's "
     "net capable power (DMN)", 75.3, "number"),
    ("pemakaian_bahan_bakar_batubara_kg", "kg", "Coal consumed", 38000000, "number"),
    ("nilai_kalor_batubara_kcal/kg", "kcal/kg", "Coal gross calorific value", 5000, "number"),
    ("pemakaian_biomassa_kg", "kg", "Biomass consumed (0 if none)", 900000, "number"),
    ("nilai_kalor_biomassa_kcal/kg", "kcal/kg", "Biomass calorific value (blank if no biomass)", 3000, "number"),
    ("pemakaian_batubara,_hsd/bio_solar,_dan_biomassa_kcal", "kcal", "Total heat input: coal + HSD/biodiesel + biomass",
     1.93e11, "number"),
    ("rencana_produksi_kwh", "kWh", "Planned gross generation", 65000000, "number"),
    ("rencana_penjualan_kwh", "kWh", "Planned net generation / sales", 58000000, "number"),
    ("rencana_pemakaian_batubara_kg", "kg", "Planned coal consumption", 40000000, "number"),
]

PRICING = [
    ("month", "date", "Last day of the month - matched to end_of_month in Data Pengusahaan", "2025-01-31", "date"),
    ("coal_price_rp_per_ton", "Rp/ton", "Coal price for that unit and month", 1850000, "number"),
    ("biomass_price_rp_per_ton", "Rp/ton", "Biomass price for that unit and month", 860000, "number"),
    ("bpp_rp_kwh", "Rp/kWh", "Biaya Pokok Penyediaan (BPP) for that unit and month - values the lost energy", 850.25, "number"),
    ("unit_name", "text", "Unit name - same spelling as 'unit' in Data Pengusahaan", "Unit 1", "text"),
]

FORMATS = {"date": "yyyy-mm-dd", "datetime": "yyyy-mm-dd hh:mm", "number": "#,##0.00", "text": "@"}


def _data_sheet(ws, required, optional=()):
    cols = [(c, True) for c in required] + [(c, False) for c in optional]
    for j, ((name, _unit, _desc, _ex, kind), req) in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=j, value=name)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = REQ_FILL if req else OPT_FILL
        cell.alignment = Alignment(vertical="center")
        letter = get_column_letter(j)
        ws.column_dimensions[letter].width = max(14, min(len(name) + 3, 48))
        for i in range(2, DATA_ROWS + 2):
            ws.cell(row=i, column=j).number_format = FORMATS[kind]
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22


def _columns_sheet(wb, title, intro, required, optional=()):
    ws = wb.create_sheet("Columns")
    ws["A1"] = title
    ws["A1"].font = Font(name=FONT, bold=True, size=13)
    for i, line in enumerate(intro, start=2):
        ws.cell(row=i, column=1, value=line).font = Font(name=FONT, color="444444")
    top = len(intro) + 3
    for j, h in enumerate(["Column", "Required", "Unit", "Description", "Example"], start=1):
        c = ws.cell(row=top, column=j, value=h)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = REQ_FILL
    rows = [(r, "Yes") for r in required] + [(r, "Optional") for r in optional]
    for i, ((name, unit, desc, ex, _kind), req) in enumerate(rows, start=top + 1):
        for j, v in enumerate([name, req, unit, desc, ex], start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = Font(name=FONT)
            c.alignment = Alignment(wrap_text=j == 4, vertical="top")
    for letter, w in zip("ABCDE", [48, 10, 11, 80, 40]):
        ws.column_dimensions[letter].width = w
    ws.freeze_panes = ws.cell(row=top + 1, column=1)


def _mapping_sheet(wb):
    """Same layout as the Status Mapping sheet read by calibrate.load_mapping (header on row 4)."""
    ws = wb.create_sheet("Status Mapping")
    ws["A1"] = "Status mapping for the outage/derating calculation"
    ws["A1"].font = Font(name=FONT, bold=True, size=13)
    ws["A2"] = ("Edit the yellow cells. Include = Yes counts the status in the model under its Calc category "
                "(FO, OS, MO, SE, FD, MD, PD). Delete this sheet to use the default mapping.")
    ws["A2"].font = Font(name=FONT, color="444444")
    head = ["status_code", "meaning", "event_status", "Calc category", "Include in calculation", "Note"]
    for j, h in enumerate(head, start=1):
        c = ws.cell(row=4, column=j, value=h)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = REQ_FILL
    for i, (code, (meaning, event)) in enumerate(STATUS_CODES.items(), start=5):
        vals = [code, meaning, event, default_category(code) or code, "Yes" if code in DEFAULT_INCLUDED else "No", None]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = Font(name=FONT)
            if j in (4, 5):
                c.fill = EDIT_FILL
    last = 4 + len(STATUS_CODES)
    dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(f"E5:E{last}")
    for letter, w in zip("ABCDEF", [12, 32, 40, 14, 22, 30]):
        ws.column_dimensions[letter].width = w
    ws.freeze_panes = "A5"
    return last


def _save(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def failure_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "All Unit"
    _data_sheet(ws, FAILURE, FAILURE_OPTIONAL)
    last = _mapping_sheet(wb)
    status_col = get_column_letter([c[0] for c in FAILURE].index("unit_status") + 1)
    dv = DataValidation(type="list", formula1=f"='Status Mapping'!$A$5:$A${last}", allow_blank=True,
                        showErrorMessage=False)
    ws.add_data_validation(dv)
    dv.add(f"{status_col}2:{status_col}{DATA_ROWS + 1}")
    _columns_sheet(wb, "Failure data template", [
        "One row per outage or derating event, all units in the 'All Unit' sheet.",
        "Dark headers are required; grey headers are optional. Extra columns are ignored.",
        "Rows split at a month boundary are merged automatically (same unit, status and cause code, gap <= 10 min).",
    ], FAILURE, FAILURE_OPTIONAL)
    return _save(wb)


def production_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    _data_sheet(ws, PRODUCTION)
    _columns_sheet(wb, "Data Pengusahaan template", [
        "One row per unit per month, in the first sheet. At least one full year per unit is recommended (the model "
        "is calibrated from annual totals; the latest year sets the most likely values).",
        "Extra columns are ignored; '-' in a cell is treated as empty.",
    ], PRODUCTION)
    return _save(wb)


def pricing_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Pricing"
    _data_sheet(ws, PRICING)
    _columns_sheet(wb, "Pricing template (optional upload)", [
        "One row per unit per month, in the first sheet. Months are matched to Data Pengusahaan by unit and "
        "calendar month.",
        "The latest year sets each unit's prices: coal weighted by coal burned, biomass by biomass burned, "
        "BPP by net sales. Without this file the model uses placeholder prices.",
    ], PRICING)
    return _save(wb)


TEMPLATES = {
    "failure": ("Template_Failure_Data.xlsx", failure_template),
    "production": ("Template_Data_Pengusahaan.xlsx", production_template),
    "pricing": ("Template_Pricing.xlsx", pricing_template),
}

assert [c[0] for c in FAILURE] == [c for c in cal.FAILURE_COLUMNS], "FAILURE template out of sync with calibrate"
assert [c[0] for c in PRODUCTION] == [c for c in cal.PRODUCTION_COLUMNS], "PRODUCTION template out of sync"
assert [c[0] for c in PRICING] == [c for c in cal.PRICING_COLUMNS], "PRICING template out of sync"


if __name__ == "__main__":
    out = Path(sys.argv[1] if len(sys.argv) > 1 else "templates")
    out.mkdir(parents=True, exist_ok=True)
    for name, make in TEMPLATES.values():
        (out / name).write_bytes(make())
        print(f"Written {out / name}")
